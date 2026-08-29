from dataclasses import dataclass
from pathlib import Path

import docx
from pypdf import PdfReader

RAW_DATA_DIR = Path("data/raw")
CATEGORIES = [
    "yonergeler",
    "usul_ve_esaslar",
    "kurum_yonetmelikleri",
    "uygulama_program_esaslari",
    "kayit_kilavuzu",
    "tip_fakultesi_mevzuati",
    "web_sss",
    "web_genel_bilgi",
]


@dataclass
class RawDocument:
    text: str
    source_file: str
    category: str


def load_pdf(path: Path) -> str:
    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def load_docx(path: Path) -> str:
    document = docx.Document(path)
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
    return "\n".join(paragraphs)


def load_document(path: Path) -> str:
    override_path = path.with_suffix(".txt")
    if override_path.exists():
        return override_path.read_text(encoding="utf-8")

    if path.suffix.lower() == ".pdf":
        return load_pdf(path)
    if path.suffix.lower() == ".docx":
        return load_docx(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def iter_corpus(raw_dir: Path = RAW_DATA_DIR):
    for category in CATEGORIES:
        category_dir = raw_dir / category
        for path in sorted(category_dir.iterdir()):
            suffix = path.suffix.lower()
            if suffix not in (".pdf", ".docx", ".txt"):
                continue
            if suffix == ".txt" and (
                path.with_suffix(".pdf").exists() or path.with_suffix(".docx").exists()
            ):
                continue  # picked up as an override by the sibling PDF/DOCX entry instead
            text = load_document(path)
            yield RawDocument(text=text, source_file=path.name, category=category)
